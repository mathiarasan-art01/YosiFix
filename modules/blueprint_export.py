"""Blueprint export: compiles the full analysis into a downloadable
Word document using python-docx."""
import io
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def build_blueprint_docx(idea, analysis, prompts=None):
    doc = Document()

    title = doc.add_heading(f"YosiFix Blueprint: {idea.title}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run(f"Domain: {idea.domain}  |  Confidence: {idea.domain_confidence*100:.0f}%  |  "
                  f"Innovation Score: {idea.innovation_score}/100").italic = True

    _heading(doc, "1. Problem Statement")
    doc.add_paragraph(idea.raw_text)

    _heading(doc, "2. Existing Solution Landscape")
    doc.add_paragraph(f"Overall similarity to closest match: {analysis.overall_similarity}%")
    doc.add_paragraph(analysis.similarity_verdict)
    for m in analysis.similar_solutions or []:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{m['name']} ({m['similarity_score']}% similar): ").bold = True
        p.add_run(m["reasoning"])

    _heading(doc, "3. Gap Analysis & Opportunities")
    doc.add_paragraph("Gaps identified:", style="Intense Quote")
    for g in analysis.gap_features or []:
        doc.add_paragraph(g, style="List Bullet")
    doc.add_paragraph("Opportunity notes:")
    for note in analysis.opportunity_notes or []:
        doc.add_paragraph(note, style="List Bullet")

    _heading(doc, "4. SDG Alignment")
    for s in analysis.sdg_mappings or []:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"SDG {s['sdg_number']}: {s['sdg_name']} — ").bold = True
        p.add_run(s["reasoning"])

    _heading(doc, "5. Recommended Tech Stack")
    stack = analysis.tech_stack or {}
    doc.add_paragraph(f"Frontend: {stack.get('frontend', '-')}")
    doc.add_paragraph(f"Backend: {stack.get('backend', '-')}")
    doc.add_paragraph(f"Database: {stack.get('database', '-')}")
    if stack.get("extra_tools"):
        doc.add_paragraph("Additional tools:")
        for t in stack["extra_tools"]:
            doc.add_paragraph(t, style="List Bullet")

    _heading(doc, "6. Build Roadmap")
    roadmap = analysis.roadmap or {}
    for phase_key, phase_label in [("mvp", "MVP (Phase 1)"), ("v2", "Version 2 (Phase 2)"), ("v3", "Version 3 (Phase 3)")]:
        doc.add_heading(phase_label, level=2)
        for item in roadmap.get(phase_key, []):
            doc.add_paragraph(item, style="List Bullet")

    if prompts:
        _heading(doc, "7. AI Build Prompt")
        p = doc.add_paragraph()
        run = p.add_run(prompts)
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
